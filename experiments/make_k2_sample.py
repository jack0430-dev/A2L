# -*- coding: utf-8 -*-
"""make_k2_sample.py — 生成 K2 历史配置样例（signal_keep_config.csv / .xlsx）

为什么需要它
------------
论文 Sec. III-B 声明知识库由三类知识源构成：K1 规则、K2 历史配置、K3 语义语料。
其中 **K2（标定表里人工标注的逐信号元数据）在代码里一直只有"加载通道"，
工作空间里却没有对应的数据文件**，因此 K2 无法被复现或验证。

本脚本把主模块内置演示数据集 ``ALL_SIGNALS`` 导出成标定表格式的历史配置，
使 K2 在工作空间中成为**可复现的数据资产**：

    data/signal_keep_config.csv    纯标准库可读（零依赖环境也能复现 K2）
    data/signal_keep_config.xlsx   与主工具的 parse_signal_config_xlsx 兼容

列集与主工具 ``SIGNAL_CONFIG_EXPORT_COLS`` 完全一致（外加 Keep 列），
因此该文件既是 RAG 知识库的 K2 来源，也能被主工具直接导入使用。

用法:
    py -3.12 experiments/make_k2_sample.py
    py -3.12 experiments/make_k2_sample.py --out-dir data --check
"""
from __future__ import annotations

import argparse
import ast as _ast
import csv
import re
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent

# 标定表列（顺序与主工具 write_signal_config_xlsx 的导出列一致）
COLUMNS = [
    "Keep", "Group", "Type", "Name", "ShortName", "DataType", "ObjectType",
    "ArraySize", "LongName", "Unit", "Min", "Max", "Factor", "Offset",
    "Description", "BitMask", "AddrOffset", "BitOffset", "BitSize",
]

# C 类型 -> A2L DATATYPE（与主模块 C_DTYPE_MAP 保持一致；此处只取标量类型）
_C_TO_A2L_DTYPE = {
    "uint8_t": "UBYTE", "int8_t": "SBYTE", "uint16_t": "UWORD",
    "int16_t": "SWORD", "uint32_t": "ULONG", "int32_t": "SLONG",
    "uint64_t": "A_UINT64", "int64_t": "A_INT64",
    "float": "FLOAT32_IEEE", "double": "FLOAT64_IEEE", "bool": "UBYTE",
}

# ALL_SIGNALS 演示行格式:
#   [Group, IO, Name, ShortName, DisplayDesc, Unit, Min, Max, Factor, Offset,
#    CType, ObjectType, Mask, Description]
_ROW_FIELDS = ("group", "io", "name", "short_name", "long_name", "unit",
               "min", "max", "factor", "offset", "c_type", "object_type",
               "bit_mask", "description")


def extract_all_signals(main_path: Path) -> list[dict]:
    """从主模块源码文本中提取 ALL_SIGNALS 的全部演示行（不 import，避免 PyQt5 依赖）。"""
    src = Path(main_path).read_text(encoding="utf-8", errors="replace")
    m = re.search(r"ALL_SIGNALS\s*=\s*\{", src)
    if not m:
        raise SystemExit("主模块中未找到 ALL_SIGNALS 定义")
    start = m.end() - 1
    depth, i = 0, start
    while i < len(src):
        if src[i] == "{":
            depth += 1
        elif src[i] == "}":
            depth -= 1
            if depth == 0:
                break
        i += 1
    block = src[start:i + 1]

    rows: list[dict] = []
    for row_m in re.finditer(r"\[[^\[\]]*\]", block):
        try:
            raw = _ast.literal_eval(row_m.group(0))
        except (ValueError, SyntaxError):
            continue
        if not (isinstance(raw, list) and len(raw) >= 14 and isinstance(raw[0], str)):
            continue
        if not str(raw[2]).strip():
            continue
        rows.append(dict(zip(_ROW_FIELDS, raw[:14])))
    return rows


def to_history_rows(all_signals: list[dict]) -> list[dict]:
    """把 ALL_SIGNALS 行转成标定表行（K2 历史配置语义）。"""
    out: list[dict] = []
    for sig in all_signals:
        c_type = str(sig["c_type"]).strip()
        row = {
            "Keep": "yes",
            "Group": str(sig["group"]).strip(),
            "Type": str(sig["io"]).strip(),          # CAL / MEAS
            "Name": str(sig["name"]).strip(),
            "ShortName": str(sig["short_name"]).strip(),
            "DataType": _C_TO_A2L_DTYPE.get(c_type, c_type),
            "ObjectType": str(sig["object_type"]).strip(),
            "ArraySize": "",
            "LongName": str(sig["long_name"]).strip(),
            "Unit": str(sig["unit"]).strip(),
            "Min": sig["min"],
            "Max": sig["max"],
            "Factor": sig["factor"],
            "Offset": sig["offset"],
            "Description": str(sig["description"]).strip(),
            "BitMask": str(sig["bit_mask"]).strip(),
            "AddrOffset": "",
            "BitOffset": "",
            "BitSize": "",
        }
        out.append(row)
    return out


def write_csv(rows: list[dict], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=COLUMNS)
        writer.writeheader()
        for row in rows:
            writer.writerow({c: row.get(c, "") for c in COLUMNS})


def write_xlsx(rows: list[dict], path: Path) -> bool:
    """写出 xlsx（需要 openpyxl）。返回是否成功。"""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        print("openpyxl 未安装，跳过 xlsx 产出（CSV 已足够复现 K2）")
        return False

    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Signals"
    header_fill = PatternFill("solid", fgColor="4472C4")
    for ci, col in enumerate(COLUMNS, start=1):
        cell = ws.cell(1, ci, col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"
    for ri, row in enumerate(rows, start=2):
        for ci, col in enumerate(COLUMNS, start=1):
            value = row.get(col, "")
            if value in ("", None):
                value = None
            ws.cell(ri, ci, value)
    ws.column_dimensions["A"].width = 8
    for ci, col in enumerate(COLUMNS, start=2):
        ws.column_dimensions[ws.cell(1, ci).column_letter].width = max(14, len(col) + 2)
    wb.save(path)
    return True


def verify(path: Path) -> int:
    """用 rag_align 的 K2 加载器回读，验证样例可被知识库消费。"""
    sys.path.insert(0, str(ROOT))
    import rag_align as ra

    rows = ra.load_historical_config_file(path)
    kb = ra.DomainKnowledgeBase(retriever=ra.CharNgramRetriever())
    added = kb.add_historical_config(rows)
    with_c_type = sum(1 for r in rows if r.get("CType"))
    print("  回读条目: %d  载入知识库: %d  带 C 类型: %d"
          % (len(rows), added, with_c_type))
    return len(rows)


def main() -> int:
    ap = argparse.ArgumentParser(description="生成 K2 历史配置样例")
    ap.add_argument("--main", default=str(ROOT / "A2L_Update_Tool_main.py"),
                    help="主模块路径（用于提取 ALL_SIGNALS）")
    ap.add_argument("--out-dir", default=str(ROOT / "data"),
                    help="输出目录（默认 data/）")
    ap.add_argument("--check", action="store_true",
                    help="生成后用 rag_align 的 K2 加载器回读校验")
    args = ap.parse_args()

    out_dir = Path(args.out_dir)
    all_signals = extract_all_signals(Path(args.main))
    rows = to_history_rows(all_signals)
    print("ALL_SIGNALS 演示信号: %d" % len(rows))

    csv_path = out_dir / "signal_keep_config.csv"
    write_csv(rows, csv_path)
    print("已写出: %s" % csv_path)

    xlsx_path = out_dir / "signal_keep_config.xlsx"
    if write_xlsx(rows, xlsx_path):
        print("已写出: %s" % xlsx_path)

    if args.check:
        for p in (csv_path, xlsx_path):
            if p.is_file():
                print("check %s" % p.name)
                verify(p)
    return 0


if __name__ == "__main__":
    sys.exit(main())
